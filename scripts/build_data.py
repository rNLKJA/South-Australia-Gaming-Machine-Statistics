# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "openpyxl>=3.1",
#   "shapely>=2.1",
#   "numpy>=2",
# ]
# ///
"""Build the website's data artefacts from the original workbook and Power BI file.

Run from the repository root:

    uv run scripts/build_data.py            # uses cached downloads when present
    uv run scripts/build_data.py --refresh  # re-download the ABS boundaries and CPI

Inputs
  original/SA Gaming Statistics.xlsx   the consolidated workbook (the source of every figure)
  original/SA Gaming Dashboard.pbix    the Power BI report (visual definitions only)
  scripts/lga_crosswalk.csv            hand-checked mapping of workbook LGA names to ABS LGAs
  ABS ASGS Edition 3 (2024) LGA boundaries for South Australia   (CC BY 4.0, no key)
  ABS Consumer Price Index, All groups, Adelaide, quarterly      (CC BY 4.0, no key)

Outputs (all derived, all small)
  web/src/data/*.json                  typed data imported by the Next.js app at build time
  web/public/data/lga-councils.geojson simplified council boundaries (71 areas)
  web/public/data/lga-groups.geojson   dissolved boundaries of every combined group CBS published
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
import sys
import urllib.parse
import urllib.request
import zipfile
from collections import OrderedDict, defaultdict
from pathlib import Path

import numpy as np
import openpyxl
import shapely
from shapely.geometry import mapping, shape

ROOT = Path(__file__).resolve().parent.parent
ORIGINAL = ROOT / "original"
WORKBOOK = ORIGINAL / "SA Gaming Statistics.xlsx"
PBIX = ORIGINAL / "SA Gaming Dashboard.pbix"
CROSSWALK = ROOT / "scripts" / "lga_crosswalk.csv"
CACHE = ROOT / "scripts" / ".cache"
OUT_DATA = ROOT / "web" / "src" / "data"
OUT_PUBLIC = ROOT / "web" / "public" / "data"

USER_AGENT = "sa-gaming-machine-stats/1.0 (+https://github.com/rNLKJA/South-Australia-Gaming-Machine-Statistics)"

ABS_LGA_URL = (
    "https://geo.abs.gov.au/arcgis/rest/services/ASGS2024/LGA/MapServer/0/query?"
    + urllib.parse.urlencode(
        {
            "where": "state_code_2021='4'",
            "outFields": "lga_code_2024,lga_name_2024,area_albers_sqkm",
            "returnGeometry": "true",
            "outSR": "4326",
            "geometryPrecision": "6",
            "f": "geojson",
        }
    )
)
ABS_CPI_URL = (
    "https://data.api.abs.gov.au/rest/data/ABS,CPI,1.1.0/1.10001.10.4.Q?startPeriod=2009-Q1"
)

# Coverage simplification tolerance in degrees (~100 m). Keeps metro councils legible.
SIMPLIFY_TOLERANCE = 0.001
COORD_PRECISION = 1e-4


# --------------------------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------------------------


def fy_key(label: str) -> str:
    """'FY 2009/10' -> '2009-10'."""
    m = re.fullmatch(r"FY (\d{4})/(\d{2})", label.strip())
    if not m:
        raise ValueError(f"unexpected financial-year label: {label!r}")
    return f"{m.group(1)}-{m.group(2)}"


def fy_of_month(year: int, month: int) -> str:
    start = year if month >= 7 else year - 1
    return f"{start}-{str(start + 1)[-2:]}"


def num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    return float(v)


def clean(v: float | None, ndigits: int = 6):
    """Round away binary noise from Excel floats and keep ints as ints."""
    if v is None:
        return None
    r = round(v, ndigits)
    return int(r) if r == int(r) else r


def sheet_rows(ws, header_row: int):
    """Yield data rows (as tuples) below the header, skipping blanks.

    The revenue sheet's used range is padded to about a million rows, so stop after a long run
    of empty rows instead of walking the whole range.
    """
    empty_run = 0
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= header_row:
            continue
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            empty_run += 1
            if empty_run > 200:
                break
            continue
        empty_run = 0
        yield row


def fetch(url: str, dest: Path, refresh: bool, accept: str | None = None) -> bytes:
    if dest.exists() and not refresh:
        return dest.read_bytes()
    dest.parent.mkdir(parents=True, exist_ok=True)
    headers = {"User-Agent": USER_AGENT}
    if accept:
        headers["Accept"] = accept
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req, timeout=120) as resp:  # noqa: S310 (fixed https URLs)
        data = resp.read()
    dest.write_bytes(data)
    return data


def write_json(path: Path, obj) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")) + "\n")
    print(f"  wrote {path.relative_to(ROOT)} ({path.stat().st_size / 1024:.1f} KB)")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# --------------------------------------------------------------------------------------------
# workbook sheets
# --------------------------------------------------------------------------------------------


def read_statewide(wb) -> list[dict]:
    ws = wb["Gaming Statistics Statewide"]
    rows = []
    for r in sheet_rows(ws, header_row=2):
        quarter, month_label, month_start, fy_label, ngr, tax, share, machines, venues = r[:9]
        d = month_start if isinstance(month_start, dt.datetime) else month_label
        fy = fy_key(fy_label)
        assert fy == fy_of_month(d.year, d.month), (fy, d)
        rows.append(
            {
                "month": f"{d.year}-{d.month:02d}",
                "fy": fy,
                "quarter": quarter,
                "ngr": clean(num(ngr), 4),
                "tax": clean(num(tax), 4),
                "venueShare": clean(num(share), 4),
                "machines": clean(num(machines)),
                "venues": clean(num(venues)),
            }
        )
    rows.sort(key=lambda x: x["month"])
    return rows


LICENCE_CATEGORIES = ["Hotels", "Clubs", "Special Circumstances", "Casino"]


def read_licences(wb) -> list[dict]:
    ws = wb["Gaming Machine Licence"]
    rows = []
    for r in sheet_rows(ws, header_row=2):
        date, fy_label, category, licences, entitlements, live_licences, live_machines = r[:7]
        assert isinstance(date, dt.datetime), r
        assert category in LICENCE_CATEGORIES, category
        fy = fy_key(fy_label)
        assert fy == fy_of_month(date.year, date.month), (fy, date)
        rows.append(
            {
                "month": f"{date.year}-{date.month:02d}",
                "fy": fy,
                "category": category,
                "licences": clean(num(licences)),
                "entitlements": clean(num(entitlements)),
                "liveLicences": clean(num(live_licences)),
                "liveMachines": clean(num(live_machines)),
            }
        )
    order = {c: i for i, c in enumerate(LICENCE_CATEGORIES)}
    rows.sort(key=lambda x: (x["month"], order[x["category"]]))
    return rows


MONTHS = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]  # fmt: skip


def read_manufacturers(wb) -> list[dict]:
    ws = next(s for s in wb.worksheets if s.title.startswith("Gaming Manufacturer"))
    rows = []
    for r in sheet_rows(ws, header_row=2):
        quarter, month_name, year, fy_label, manufacturer, machines, pct, total = r[:8]
        m = MONTHS.index(month_name.strip()) + 1
        y = int(year)
        fy = fy_key(fy_label)
        assert fy == fy_of_month(y, m), (fy, y, m)
        rows.append(
            {
                "month": f"{y}-{m:02d}",
                "fy": fy,
                "quarter": quarter,
                "manufacturer": manufacturer.strip(),
                "machines": clean(num(machines)),
                "pct": clean(num(pct), 6),
                "monthTotal": clean(num(total)),
            }
        )
    rows.sort(key=lambda x: (x["month"], -x["machines"]))
    return rows


def read_lga(wb) -> list[dict]:
    ws = wb["Gaming Machine Revenue"]
    rows = []
    for r in sheet_rows(ws, header_row=2):
        name, ngr, avg, machines, premises, fy_label = r[:6]
        rows.append(
            {
                "fy": fy_key(fy_label),
                "name": str(name).strip(),
                "ngr": clean(num(ngr), 4),
                "avgPerVenue": clean(num(avg), 4),
                "machines": clean(num(machines), 4),
                "premises": clean(num(premises), 4),
            }
        )
    return rows


def read_info_pivots(wb) -> dict:
    """Pull the 'Average of ...' / 'Sum of ...' pivot tables at the foot of the INFO sheet.

    These are the checks the original author built next to the Power BI report; the web app's
    unit tests compare its own aggregations against them (parity tests).
    """
    ws = wb["INFO"]
    lines = []
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row[1:] if v is not None]
        if vals:
            lines.append(list(row[1:]))
    pivots: dict[str, dict] = {}
    i = 0
    while i < len(lines):
        title = lines[i][0]
        if isinstance(title, str) and (title.startswith("Average of") or title.startswith("Sum of")):
            # header row follows (category label + FY columns)
            header = lines[i + 1]
            fys = [h for h in header[1:] if isinstance(h, str) and h.startswith("FY")]
            table: dict[str, dict[str, float]] = {}
            j = i + 2
            while j < len(lines):
                label = lines[j][0]
                if not isinstance(label, str) or label.startswith(("Average of", "Sum of")):
                    break
                if label in ("Grand Total",) or label in LICENCE_CATEGORIES or label.startswith("FY"):
                    values = lines[j][1 : 1 + len(fys)]
                    table[label] = {fy_key(f): v for f, v in zip(fys, values) if v is not None}
                    j += 1
                    continue
                break
            if table:
                pivots[title] = table
            i = j
            continue
        if title == "Values":  # statewide 'Sum of ...' block: FY columns, measures as rows
            fys = [h for h in lines[i][1:] if isinstance(h, str) and h.startswith("FY")]
            j = i + 1
            while j < len(lines) and isinstance(lines[j][0], str) and lines[j][0].startswith("Sum of"):
                values = lines[j][1 : 1 + len(fys)]
                pivots[f"Statewide: {lines[j][0]}"] = {
                    "all": {fy_key(f): v for f, v in zip(fys, values) if v is not None}
                }
                j += 1
            i = j
            continue
        if title == "Financial Year" and i > 0 and lines[i - 1][0] == "Average of No. of GMs":
            pass
        i += 1

    # Manufacturer pivot: FY rows x manufacturer columns.
    for k, line in enumerate(lines):
        if line[0] == "Average of No. of GMs":
            header = lines[k + 1]
            makers = [h for h in header[1:] if isinstance(h, str) and h != "Grand Total"]
            table = defaultdict(dict)
            j = k + 2
            while j < len(lines) and isinstance(lines[j][0], str) and lines[j][0].startswith("FY"):
                fy = fy_key(lines[j][0])
                for mk, v in zip(makers, lines[j][1:]):
                    if v is not None:
                        table[mk][fy] = v
                j += 1
            pivots["Average of No. of GMs"] = dict(table)
        if line[0] == "Sum of Aggregated NGR by LGA (AUD)":
            header = lines[k + 1]
            fys = [h for h in header[1:] if isinstance(h, str) and h.startswith("FY")]
            table = {}
            j = k + 2
            while j < len(lines) and isinstance(lines[j][0], str):
                label = lines[j][0]
                table[label] = {fy_key(f): v for f, v in zip(fys, lines[j][1:]) if v is not None}
                if label == "Grand Total":
                    break
                j += 1
            pivots["Sum of Aggregated NGR by LGA (AUD)"] = table
    return pivots


# --------------------------------------------------------------------------------------------
# LGA: crosswalk + reconstruction of the units CBS actually published
# --------------------------------------------------------------------------------------------


def read_crosswalk() -> dict[str, dict]:
    with CROSSWALK.open(newline="") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        if None in r or any(v is None for v in r.values()):
            raise SystemExit(f"crosswalk row has the wrong number of fields (quote commas): {r}")
    return {r["workbook_name"]: r for r in rows}


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def build_units(lga_rows: list[dict], crosswalk: dict[str, dict], abs_codes: dict[str, str]):
    """Rebuild the rows CBS published from the workbook's per-name rows.

    The workbook repeats a combined group's figures across every member name after dividing them
    equally, so members of one published row share an identical (NGR, NGR per venue) signature
    within a financial year. (Machines and premises are not used in the signature: when a group's
    count is odd the workbook rounds the halves, e.g. Light 54 + Mallala 55 machines in FY 2013/14.)
    Group on that signature, then map every name through the crosswalk. A unit with one display name is a single council (possibly split into
    fragments by the workbook); a unit with several is a combined group.
    """
    by_fy: dict[str, list[dict]] = defaultdict(list)
    for r in lga_rows:
        by_fy[r["fy"]].append(r)

    units = []
    for fy in sorted(by_fy):
        rows = by_fy[fy]
        machines_published = any((r["machines"] or 0) > 0 for r in rows)
        buckets: OrderedDict[tuple, list[dict]] = OrderedDict()
        for r in rows:
            sig = (r["ngr"], r["avgPerVenue"])
            buckets.setdefault(sig, []).append(r)
        for sig, members in buckets.items():
            names = [m["name"] for m in members]
            missing = [n for n in names if n not in crosswalk]
            if missing:
                raise SystemExit(f"crosswalk is missing workbook names: {missing}")
            display = list(OrderedDict.fromkeys(crosswalk[n]["display_name"] for n in names))
            display_sorted = sorted(display)
            geo_lgas = sorted(
                {crosswalk[n]["abs_lga_2024"] for n in names if crosswalk[n]["geometry"] == "yes"}
            )
            codes = sorted(abs_codes[a] for a in geo_lgas)
            n = len(members)
            ngr_total = round(sum(m["ngr"] for m in members), 2)
            premises_total = round(sum(m["premises"] for m in members))
            machines_total = (
                round(sum(m["machines"] for m in members)) if machines_published else None
            )
            relations = sorted({crosswalk[nm]["relation"] for nm in names})
            units.append(
                {
                    "fy": fy,
                    "id": slug(" + ".join(display_sorted)),
                    "label": ", ".join(display_sorted),
                    "members": display_sorted,
                    "kind": "group" if len(display) > 1 else ("split" if n > 1 else "single"),
                    "workbookNames": names,
                    "workbookRows": n,
                    "codes": codes,
                    "geoKey": "+".join(codes),
                    "ngr": ngr_total,
                    "avgPerVenue": members[0]["avgPerVenue"],
                    "machines": machines_total,
                    "premises": premises_total,
                    "perRowNgr": members[0]["ngr"],
                    "relations": relations,
                }
            )
    return units


# --------------------------------------------------------------------------------------------
# geometry
# --------------------------------------------------------------------------------------------


def build_geometry(raw: bytes, units: list[dict]):
    fc = json.loads(raw)
    feats = [f for f in fc["features"] if f.get("geometry")]
    codes = [f["properties"]["lga_code_2024"] for f in feats]
    names = [f["properties"]["lga_name_2024"] for f in feats]
    areas = [f["properties"].get("area_albers_sqkm") for f in feats]
    geoms = np.array([shape(f["geometry"]) for f in feats], dtype=object)
    if not shapely.coverage_is_valid(geoms):
        raise SystemExit("ABS LGA coverage is not valid; cannot simplify safely")
    simp = shapely.set_precision(
        shapely.coverage_simplify(geoms, tolerance=SIMPLIFY_TOLERANCE), COORD_PRECISION
    )
    by_code = dict(zip(codes, simp))

    def feature(geom, props):
        geom = shapely.make_valid(geom)
        return {"type": "Feature", "properties": props, "geometry": mapping(geom)}

    councils_fc = {
        "type": "FeatureCollection",
        "features": [
            feature(by_code[c], {"code": c, "name": n, "areaSqKm": round(a or 0, 1)})
            for c, n, a in zip(codes, names, areas)
        ],
    }

    group_keys = OrderedDict()
    for u in units:
        if len(u["codes"]) > 1:
            group_keys.setdefault(u["geoKey"], u["codes"])
    groups_fc = {
        "type": "FeatureCollection",
        "features": [
            feature(
                shapely.set_precision(shapely.union_all([by_code[c] for c in cs]), COORD_PRECISION),
                {"geoKey": k, "codes": cs},
            )
            for k, cs in group_keys.items()
        ],
    }
    councils = [
        {"code": c, "name": n, "areaSqKm": round(a or 0, 1)} for c, n, a in zip(codes, names, areas)
    ]
    councils.sort(key=lambda x: x["name"])
    return councils, councils_fc, groups_fc


def round_coords(obj, nd=4):
    if isinstance(obj, float):
        return round(obj, nd)
    if isinstance(obj, list):
        return [round_coords(x, nd) for x in obj]
    if isinstance(obj, tuple):
        return [round_coords(x, nd) for x in obj]
    if isinstance(obj, dict):
        return {k: round_coords(v, nd) for k, v in obj.items()}
    return obj


# --------------------------------------------------------------------------------------------
# CPI
# --------------------------------------------------------------------------------------------


def build_cpi(raw: bytes) -> list[dict]:
    text = raw.decode("utf-8")
    reader = csv.DictReader(text.splitlines())
    series = []
    for r in reader:
        period = r.get("TIME_PERIOD") or r.get("TIME_PERIOD: Time Period")
        value = r.get("OBS_VALUE")
        if period and value:
            series.append({"quarter": period.split(":")[0].strip(), "index": float(value)})
    series.sort(key=lambda x: x["quarter"])
    if not series:
        raise SystemExit("ABS CPI response had no observations")
    return series


# --------------------------------------------------------------------------------------------
# Power BI report definition
# --------------------------------------------------------------------------------------------


def read_powerbi() -> list[dict]:
    with zipfile.ZipFile(PBIX) as z:
        layout = json.loads(z.read("Report/Layout").decode("utf-16-le"))
    visuals = []
    for section in layout["sections"]:
        for vc in section.get("visualContainers", []):
            cfg = json.loads(vc["config"])
            sv = cfg.get("singleVisual", {})
            projections = {
                role: [p["queryRef"] for p in items]
                for role, items in sv.get("projections", {}).items()
            }
            if not projections:
                continue
            visuals.append(
                {
                    "page": section["displayName"],
                    "visualType": sv.get("visualType"),
                    "projections": projections,
                }
            )
    return visuals


# --------------------------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------------------------


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--refresh", action="store_true", help="re-download ABS boundaries and CPI")
    args = ap.parse_args()

    print("reading workbook …")
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    statewide = read_statewide(wb)
    licences = read_licences(wb)
    manufacturers = read_manufacturers(wb)
    lga_rows = read_lga(wb)
    wb.close()
    # INFO has merged cells; read it with the normal (non-streaming) loader.
    wb_full = openpyxl.load_workbook(WORKBOOK, data_only=True)
    pivots = read_info_pivots(wb_full)
    wb_full.close()
    print(
        f"  statewide {len(statewide)} · licences {len(licences)} · "
        f"manufacturers {len(manufacturers)} · lga {len(lga_rows)} rows"
    )
    assert len(statewide) == 180 and len(licences) == 642
    assert len(manufacturers) == 1633 and len(lga_rows) == 913

    print("fetching ABS LGA 2024 boundaries and Adelaide CPI …")
    lga_raw = fetch(ABS_LGA_URL, CACHE / "abs_lga_2024_sa.geojson", args.refresh)
    cpi_raw = fetch(
        ABS_CPI_URL,
        CACHE / "abs_cpi_adelaide.csv",
        args.refresh,
        accept="application/vnd.sdmx.data+csv;labels=id",
    )
    abs_codes = {
        f["properties"]["lga_name_2024"]: f["properties"]["lga_code_2024"]
        for f in json.loads(lga_raw)["features"]
    }

    crosswalk = read_crosswalk()
    unknown = sorted({r["abs_lga_2024"] for r in crosswalk.values()} - abs_codes.keys())
    if unknown:
        raise SystemExit(f"crosswalk names not found in ABS LGA 2024: {unknown}")

    units = build_units(lga_rows, crosswalk, abs_codes)
    councils, councils_fc, groups_fc = build_geometry(lga_raw, units)

    # ---- sanity checks that would make the site wrong if they failed ----
    totals = defaultdict(float)
    for r in lga_rows:
        totals[r["fy"]] += r["ngr"]
    unit_totals = defaultdict(float)
    for u in units:
        unit_totals[u["fy"]] += u["ngr"]
    for fy in totals:
        assert abs(totals[fy] - unit_totals[fy]) < 1.0, (fy, totals[fy], unit_totals[fy])
    # No ABS polygon may be claimed by two units in the same year.
    for fy in sorted({u["fy"] for u in units}):
        seen: dict[str, str] = {}
        for u in (u for u in units if u["fy"] == fy):
            for c in u["codes"]:
                if c in seen:
                    raise SystemExit(f"{fy}: {c} in both {seen[c]} and {u['label']}")
                seen[c] = u["label"]

    cpi = build_cpi(cpi_raw)
    powerbi = read_powerbi()

    print("writing artefacts …")
    write_json(
        OUT_DATA / "statewide.json",
        {"source": "original/SA Gaming Statistics.xlsx · Gaming Statistics Statewide", "rows": statewide},
    )
    write_json(
        OUT_DATA / "licences.json",
        {"source": "original/SA Gaming Statistics.xlsx · Gaming Machine Licence", "rows": licences},
    )
    write_json(
        OUT_DATA / "manufacturers.json",
        {"source": "original/SA Gaming Statistics.xlsx · Gaming Manufacturer’s Market", "rows": manufacturers},
    )
    crosswalk_out = [
        {
            "workbookName": r["workbook_name"],
            "displayName": r["display_name"],
            "absName": r["abs_lga_2024"],
            "absCode": abs_codes[r["abs_lga_2024"]],
            "relation": r["relation"],
            "geometry": r["geometry"] == "yes",
            "note": r["note"],
            "years": sorted({x["fy"] for x in lga_rows if x["name"] == r["workbook_name"]}),
        }
        for r in crosswalk.values()
    ]
    write_json(
        OUT_DATA / "lga.json",
        {
            "source": "original/SA Gaming Statistics.xlsx · Gaming Machine Revenue",
            "rows": lga_rows,
            "units": units,
            "crosswalk": crosswalk_out,
            "councils": councils,
        },
    )
    write_json(
        OUT_DATA / "cpi.json",
        {
            "source": "ABS Consumer Price Index, Australia (CPI 1.1.0): index numbers, All groups CPI, Adelaide, original, quarterly",
            "url": ABS_CPI_URL,
            "series": cpi,
        },
    )
    write_json(OUT_DATA / "powerbi.json", {"source": "original/SA Gaming Dashboard.pbix · Report/Layout", "visuals": powerbi})
    write_json(OUT_DATA / "info-pivots.json", {"source": "original/SA Gaming Statistics.xlsx · INFO", "pivots": pivots})
    write_json(
        OUT_DATA / "meta.json",
        {
            "workbook": {"path": "original/SA Gaming Statistics.xlsx", "sha256": sha256(WORKBOOK)},
            "powerbi": {"path": "original/SA Gaming Dashboard.pbix", "sha256": sha256(PBIX)},
            "boundaries": {
                "name": "ABS Australian Statistical Geography Standard (ASGS) Edition 3, Local Government Areas 2024, South Australia",
                "licence": "CC BY 4.0",
                "url": "https://geo.abs.gov.au/arcgis/rest/services/ASGS2024/LGA/MapServer",
                "simplifyToleranceDeg": SIMPLIFY_TOLERANCE,
            },
            "cpiLatestQuarter": cpi[-1]["quarter"],
        },
    )
    OUT_PUBLIC.mkdir(parents=True, exist_ok=True)
    for name, fc in (("lga-councils.geojson", councils_fc), ("lga-groups.geojson", groups_fc)):
        path = OUT_PUBLIC / name
        path.write_text(json.dumps(round_coords(fc), separators=(",", ":")) + "\n")
        print(f"  wrote {path.relative_to(ROOT)} ({path.stat().st_size / 1024:.1f} KB)")

    groups = [u for u in units if u["kind"] == "group"]
    splits = [u for u in units if u["kind"] == "split"]
    print(
        f"done: {len(units)} published LGA rows ({len(groups)} combined groups, "
        f"{len(splits)} single councils split by name in the workbook), "
        f"{len(groups_fc['features'])} distinct group shapes, CPI to {cpi[-1]['quarter']}"
    )


if __name__ == "__main__":
    sys.exit(main())
